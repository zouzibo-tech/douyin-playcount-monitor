"""Excel 日报导出：Sheet1 合集总览，Sheet2 分集明细，Sheet3 账号汇总。"""
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import analyze
from config import load_config, resolve_output_dir


def _style_header(ws, ncols):
    fill = PatternFill("solid", fgColor="1F4E79")
    thin = Side(style="thin", color="BFBFBF")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws.freeze_panes = "A2"


def _style_body(ws, ncols, nrows):
    thin = Side(style="thin", color="D9D9D9")
    for r in range(2, nrows + 2):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="微软雅黑", size=10)
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            cell.alignment = Alignment(vertical="center")


def _autofit(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def export_daily(cfg=None):
    cfg = cfg or load_config()
    ov = analyze.overview()
    out_dir = resolve_output_dir(cfg)
    out_dir.mkdir(parents=True, exist_ok=True)

    fd = ov.get("fetch_date") or datetime.now().strftime("%Y-%m-%d")
    path = out_dir / f"播放量日报_{fd}.xlsx"

    wb = Workbook()

    # ---- Sheet1 合集总览 ----
    ws = wb.active
    ws.title = "合集总览"
    headers = ["序号", "合集名", "账号", "总播放量", "昨日总播放量",
               "日增", "收藏量", "集数", "更新至", "分集求和(校验)", "数据日期"]
    ws.append(headers)
    prev_play = {}
    for r in ov["mixes"]:
        prev_play[r["mix_id"]] = (r.get("play_vv") or 0) - (r.get("delta") or 0) \
            if r.get("delta") is not None else None
    for i, r in enumerate(ov["mixes"], 1):
        p = prev_play.get(r["mix_id"])
        ws.append([
            i, r.get("mix_name"), r.get("account"),
            r.get("play_vv"), p if p is not None else "—",
            r.get("delta") if r.get("delta") is not None else "—",
            r.get("collect_vv"), r.get("episode_count"),
            r.get("updated_to_episode"), r.get("play_sum_calc"), fd,
        ])
    _style_header(ws, len(headers))
    _style_body(ws, len(headers), len(ov["mixes"]))
    _autofit(ws, [6, 40, 16, 14, 14, 12, 10, 8, 8, 16, 12])

    # ---- Sheet2 分集明细 ----
    ws2 = wb.create_sheet("分集明细")
    h2 = ["合集名", "账号", "集数", "标题", "播放量", "昨日播放量",
          "日增", "占合集比", "点赞", "评论", "分享", "收藏"]
    ws2.append(h2)
    row_n = 0
    for r in ov["mixes"]:
        det = analyze.mix_detail(r["mix_id"], r)
        for e in det["episodes"]:
            pc = e.get("play_count") or 0
            d = e.get("delta")
            ws2.append([
                r.get("mix_name"), r.get("account"), f"第{e.get('episode_no')}集",
                e.get("title"), pc,
                (pc - d) if d is not None else "—",
                d if d is not None else "—",
                f"{e.get('share')}%" if e.get("share") is not None else "—",
                e.get("digg_count"), e.get("comment_count"),
                e.get("share_count"), e.get("collect_count"),
            ])
            row_n += 1
    _style_header(ws2, len(h2))
    _style_body(ws2, len(h2), row_n)
    _autofit(ws2, [36, 16, 8, 46, 12, 12, 12, 10, 10, 10, 10, 10])

    # ---- Sheet3 账号汇总 ----
    ws3 = wb.create_sheet("账号汇总")
    h3 = ["账号", "合集数", "总播放量", "昨日增量"]
    ws3.append(h3)
    for a in ov["accounts"]:
        ws3.append([a["account"], a["mixes"], a["play"], a["delta"]])
    _style_header(ws3, len(h3))
    _style_body(ws3, len(h3), len(ov["accounts"]))
    _autofit(ws3, [24, 10, 16, 14])

    wb.save(str(path))
    return str(path)
